#!/usr/bin/env python3
"""Maakt Word-bestanden (.docx) van de modellen in modellen.py.

Gebruik:
  python3 build/modellen_docx.py <map>            alle modellen (ook op aanvraag en verborgen)
  build.py roept make_docx() zelf aan voor de modellen met status "download".
"""
import sys
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.enum.text import WD_COLOR_INDEX, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

sys.path.insert(0, str(Path(__file__).parent))
from modellen import MODELLEN, ALGEMEEN, VERSIE, GEBIEDEN

NAVY = RGBColor(0x1B, 0x3D, 0x5C)
GRIJS = RGBColor(0x5B, 0x6B, 0x7B)
LOGO = Path(__file__).parent / "assets" / "logo-mark.png"
GEBIED_NAAM = dict(GEBIEDEN)


def runs(par, text, size=None, color=None, bold=False, italic=False):
    """Zet tekst in een alinea; alles tussen [ ] krijgt een lichte markering."""
    depth, buf = 0, ""
    def flush(highlight):
        nonlocal buf
        if not buf:
            return
        r = par.add_run(buf)
        r.bold, r.italic = bold, italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        if highlight: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        buf = ""
    for ch in text:
        if ch == "[":
            if depth == 0: flush(False)
            depth += 1; buf += ch
        elif ch == "]" and depth:
            buf += ch; depth -= 1
            if depth == 0: flush(True)
        else:
            buf += ch
    flush(depth > 0)


def lines(par, text, **kw):
    parts = text.split("\n")
    for i, part in enumerate(parts):
        runs(par, part, **kw)
        if i < len(parts) - 1:
            par.add_run().add_break(WD_BREAK.LINE)


def shade(par, fill="EEF3F8"):
    ppr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def page_number(par):
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        r = par.add_run()
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
        r._r.append(el)
        r.font.size = Pt(8); r.font.color.rgb = GRIJS


def blocks(doc, items, tint=False):
    for b in items:
        if isinstance(b, str):
            p = doc.add_paragraph(); lines(p, b, bold=b.startswith("Betreft:"))
        elif b[0] == "h":
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.keep_with_next = True
            runs(p, b[1], size=11.5, color=NAVY, bold=True)
        elif b[0] in ("li", "ol"):
            for i, item in enumerate(b[1]):
                if b[0] == "ol":
                    # eigen nummering per lijst, zodat elke lijst weer bij 1 begint
                    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.first_line_indent = Cm(-0.6)
                    runs(p, f"{i + 1}. "); runs(p, item)
                else:
                    p = doc.add_paragraph(style="List Bullet"); runs(p, item)
                if tint: shade(p)
            continue
        elif b[0] == "klein":
            p = doc.add_paragraph(); runs(p, b[1], size=9, color=GRIJS)
        elif b[0] == "onder":
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(18)
            lines(p, "\n".join(b[1]))
        else:
            raise ValueError(b)
        if tint: shade(p)


def make_docx(m, path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.3); sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.0)
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15

    # koptekst: logo en naam
    hp = sec.header.paragraphs[0]
    hp.add_run().add_picture(str(LOGO), height=Cm(0.9))
    r = hp.add_run("   LAW & RESEARCH"); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = NAVY
    # voettekst
    fp = sec.footer.paragraphs[0]
    runs(fp, f"Model van Law & Research · law-research.nl · versie {VERSIE} · algemene informatie, geen juridisch advies · pagina ", size=8, color=GRIJS)
    page_number(fp)

    p = doc.add_paragraph(); runs(p, f"{GEBIED_NAAM[m['gebied']].upper()} · {m['soort'].upper()}", size=8.5, color=GRIJS, bold=True)
    p = doc.add_paragraph(); runs(p, m["titel"], size=17, color=NAVY, bold=True)
    p = doc.add_paragraph(); runs(p, m["kort"], italic=True, color=GRIJS)

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); shade(p)
    runs(p, "Toelichting  ", bold=True, color=NAVY); runs(p, "(haal dit deel weg voordat u het model gebruikt)", size=9, color=GRIJS)
    blocks(doc, m["toelichting"], tint=True)
    p = doc.add_paragraph(); shade(p); runs(p, ALGEMEEN, size=9, color=GRIJS)
    p = doc.add_paragraph(); runs(p, "Geel gemarkeerde tekst tussen [vierkante haken] vult u zelf in of schrapt u.", size=9, color=GRIJS)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    blocks(doc, m["tekst"])

    core = doc.core_properties
    core.title, core.author, core.subject = m["titel"], "Law & Research", GEBIED_NAAM[m["gebied"]]
    core.created = core.modified = datetime(2026, 10, 30)
    core.last_modified_by = "Law & Research"
    core.comments = "Model van Law & Research (law-research.nl), versie " + VERSIE
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def bestandsnaam(m):
    return f"{m['slug']}.docx"


if __name__ == "__main__":
    out = Path(sys.argv[1])
    for m in MODELLEN:
        make_docx(m, out / m["gebied"] / bestandsnaam(m))
    print(len(MODELLEN), "modellen")
